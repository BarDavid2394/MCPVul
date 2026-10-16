"""Build the editable project book and figures from the Markdown source.

Requires python-docx and matplotlib. Document dependencies may be isolated in
%TEMP%/mcp-book-libs; the scanner itself gains no dependency.
"""
from pathlib import Path
import json
import os
import re
import sys
import tempfile

libs = Path(tempfile.gettempdir())/'mcp-book-libs'
if libs.exists():
    sys.path.insert(0, str(libs))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ASSETS = HERE/'assets'
ASSETS.mkdir(exist_ok=True)
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,
                     'axes.spines.right':False,'savefig.dpi':220})
ink = '#24384b'
accent = '#52728b'


def figures():
    fig, ax = plt.subplots(figsize=(8,5.6))
    ax.set(xlim=(0,10),ylim=(0,7));ax.axis('off')
    boxes = [(.4,5.8,9.2,.7,'Source code • MCP configuration • Dependency manifests'),
             (.4,4.35,2.5,.9,'Artifact\ndiscovery'),(3.75,4.35,2.5,.9,'Python / TS / JS\nparsing'),
             (7.1,4.35,2.5,.9,'Project context\n& bounded flow'),
             (3.1,2.8,3.8,.9,'Catalog analyzers\nMCP01–MCP16'),
             (3.1,1.25,3.8,.9,'Evidence, filtering\n& deduplication'),
             (.25,.05,2.6,.75,'Console / JSON\nreports'),(3.7,.05,2.6,.75,'Optional advisory\nreview'),
             (7.15,.05,2.6,.75,'Eligible MCP03/06\nremediation')]
    for x,y,w,h,t in boxes:
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.05',facecolor='#f2f5f7',edgecolor=accent,lw=1.2))
        ax.text(x+w/2,y+h/2,t,ha='center',va='center',color=ink,fontsize=10)
    for a,b in [((5,5.8),(1.65,5.3)),((2.95,4.8),(3.7,4.8)),((6.3,4.8),(7.05,4.8)),
                ((8.35,4.3),(6,3.75)),((5,4.3),(5,3.75)),((5,2.75),(5,2.2)),
                ((4,1.2),(1.55,.85)),((5,1.2),(5,.85)),((6,1.2),(8.45,.85))]:
        ax.add_patch(FancyArrowPatch(a,b,arrowstyle='-|>',mutation_scale=11,color=accent,lw=1.2))
    fig.tight_layout();fig.savefig(ASSETS/'architecture.png',bbox_inches='tight');fig.savefig(ASSETS/'architecture.svg',bbox_inches='tight');plt.close(fig)

    fig,ax=plt.subplots(figsize=(7.6,2.15))
    labels=['Credible evidence','False positive','Needs context'];values=[13,32,38]
    bars=ax.barh(labels,values,color=['#4e7864','#ad665b','#8a929a'],height=.55)
    ax.bar_label(bars,padding=5);ax.set_xlim(0,44);ax.set_xlabel('Reviewed findings (historical run, n = 83)');ax.invert_yaxis()
    fig.tight_layout();fig.savefig(ASSETS/'finding_labels.png',bbox_inches='tight');fig.savefig(ASSETS/'finding_labels.svg',bbox_inches='tight');plt.close(fig)

    fig,ax=plt.subplots(figsize=(7.6,1.8))
    ax.barh(['485 source files'],[393],color=accent,label='Flagged: 393 (81.0%)')
    ax.barh(['485 source files'],[92],left=[393],color='#d4dbe0',label='Unflagged: 92 (19.0%)')
    ax.text(196.5,0,'393 flagged',ha='center',va='center',color='white',weight='bold')
    ax.text(439,0,'92',ha='center',va='center',color=ink)
    ax.set_xlim(0,485);ax.set_xlabel('Exposed development corpus; not an independent holdout')
    fig.tight_layout();fig.savefig(ASSETS/'mcptox.png',bbox_inches='tight');fig.savefig(ASSETS/'mcptox.svg',bbox_inches='tight');plt.close(fig)

    d=json.loads((HERE/'data_audit.json').read_text())['holdout']['rule_counts']
    labels=[k.replace('MCP','').replace('UNTRUSTED-CONTEXT','CONTEXT').replace('COMMAND-INJECTION','COMMAND') for k in sorted(d)]
    pos=[d[k].get('positive',0) for k in sorted(d)];neg=[d[k].get('negative',0) for k in sorted(d)]
    fig,ax=plt.subplots(figsize=(7.6,3.35));yy=list(range(len(labels)))
    ax.barh([y-.18 for y in yy],pos,height=.33,color=accent,label='Positive')
    ax.barh([y+.18 for y in yy],neg,height=.33,color='#a6b5c1',label='Negative')
    ax.set_yticks(yy,labels);ax.invert_yaxis();ax.set_xlim(0,6);ax.set_xticks(range(6));ax.set_xlabel('Explicit artifact/rule labels');ax.legend(loc='lower right',frameon=False)
    fig.tight_layout();fig.savefig(ASSETS/'holdout_coverage.png',bbox_inches='tight');fig.savefig(ASSETS/'holdout_coverage.svg',bbox_inches='tight');plt.close(fig)


def shade(cell, color):
    el=OxmlElement('w:shd');el.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(el)


def make_table(doc, lines):
    rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[-: ]+',v) for v in r)]
    tab=doc.add_table(rows=0,cols=len(rows[0]));tab.style='Table Grid'
    for i, vals in enumerate(rows):
        cells=tab.add_row().cells
        for cell,val in zip(cells,vals):
            cell.text=val
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                p.paragraph_format.line_spacing=1.0
                for run in p.runs:run.font.size=Pt(9.5);run.bold=i==0
            if i==0:shade(cell,'EDF0F2')
        trPr=tab.rows[-1]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
    hdr=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(hdr)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def main():
    figures()
    pages=(HERE/'MCP_Project_Book.md').read_text(encoding='utf-8').split('<!-- PAGE -->')
    assert len(pages)==34,len(pages)
    doc=Document();s=doc.sections[0]
    s.page_width=Cm(21);s.page_height=Cm(29.7)
    s.left_margin=Cm(2.54);s.right_margin=Cm(2.54);s.top_margin=Cm(2.3);s.bottom_margin=Cm(2.0)
    s.footer_distance=Cm(1.0)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(11)
    normal.paragraph_format.line_spacing=1.1;normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',27),('Heading 1',25),('Heading 2',13)]:
        style=doc.styles[name];style.font.name='Calibri Light' if size>20 else 'Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(12 if size>20 else 5)
        style.paragraph_format.space_before=Pt(0 if size>20 else 8)
    doc.styles['Heading 2'].font.bold=True
    foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    doc.core_properties.title='Static Security Analysis of Model Context Protocol Servers'
    doc.core_properties.author='Bar David and Noam Ben Lulu'
    doc.core_properties.subject='B.Sc. final project, Technion; validation snapshot 11 October 2026'
    for page_num,page in enumerate(pages,1):
        if page_num>1:doc.add_page_break()
        lines=page.strip().splitlines();i=0
        if page_num==1:
            for line in lines:
                if not line.strip():continue
                if line.startswith('# '):
                    p=doc.add_paragraph(line[2:],style='Title');p.paragraph_format.space_before=Pt(42);p.paragraph_format.space_after=Pt(24)
                else:
                    p=doc.add_paragraph(line);p.paragraph_format.space_after=Pt(10)
                    if line in ['Submitted by','Supervisor']:p.paragraph_format.space_before=Pt(16);p.runs[0].bold=True
                    if line.startswith(('Submission date','Implementation')):
                        for run in p.runs:run.font.size=Pt(10)
                    if re.search(r'[\u0590-\u05ff]',line):
                        for run in p.runs:
                            rtl=OxmlElement('w:rtl');run._r.get_or_add_rPr().append(rtl)
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            continue
        while i<len(lines):
            line=lines[i].strip()
            if not line:i+=1;continue
            if line.startswith('|'):
                chunk=[]
                while i<len(lines) and lines[i].strip().startswith('|'):chunk.append(lines[i]);i+=1
                make_table(doc,chunk);continue
            if line.startswith('```'):
                code=[];i+=1
                while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
                p=doc.add_paragraph();p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(9)
                r=p.add_run('\n'.join(code));r.font.name='Consolas';r.font.size=Pt(9)
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'F5F5F5');p._p.get_or_add_pPr().append(shd)
                i+=1;continue
            image_match=re.match(r'!\[(.*?)\]\((.*?)\)',line)
            if image_match:
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(HERE/image_match.group(2)),width=Inches(6.05))
                i+=1;continue
            if line.startswith('# '):p=doc.add_paragraph(line[2:],style='Heading 1')
            elif line.startswith('## '):p=doc.add_paragraph(line[3:],style='Heading 2')
            else:
                p=doc.add_paragraph(line)
                if line.startswith('Figure '):
                    for r in p.runs:r.italic=True;r.font.size=Pt(9)
                    p.paragraph_format.space_after=Pt(9)
                if page_num==34:
                    p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1.0
                    for r in p.runs:r.font.size=Pt(9.5)
            i+=1
    path=HERE/'MCP_Project_Book.docx';doc.save(path)
    print('Saved',path,'planned pages',len(pages),'words',len(re.findall(r'\b\S+\b',' '.join(pages))))


if __name__=='__main__':main()
